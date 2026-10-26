"""JSON-valued spreadsheet cells preserve types and avoid formula interpretation."""
from __future__ import annotations
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sqlite3
from urllib.parse import urlsplit
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from .core import CollectorError, canonical, load_records, job_lock, status, get_meta

OVERFLOW = '@sionlab-overflow:'

def columns_for(rows):
    names = list(dict.fromkeys(name for row in rows for name in row['fields']))
    columns = [('identity', ['identity']), ('source_url', ['source_url'])]
    for index, name in enumerate(names, 1):
        for key in ('raw', 'normalized', 'missing_reason'):
            # Index is authoritative: arbitrary customer field names cannot collide.
            label = re.sub(r'[\x00-\x1f]', '_', name)[:80]
            columns.append((f'f{index}:{label}:{key}', ['fields', name, key]))
    return columns

def resolve(row, path):
    result = row
    for key in path:
        if not isinstance(result, dict):
            return None
        result = result.get(key)
    return result

def safe_json(value):
    # Keep readable Korean while escaping XML-invalid noncharacters/surrogates.
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return re.sub(r'[\ud800-\udfff\ufffe\uffff]', lambda match: '\\u' + format(ord(match.group()), '04x'), encoded)

def sheet_name(name):
    clean = re.sub(r'[\\/*?:\[\]\x00-\x1f]', '_', name).strip("'")[:31] or 'Data'
    if clean.lower() in ('_overflow', '_columns', 'overview', 'collection'):
        clean = clean[:26] + '_data'
    return clean

def string_cells(sheet, values):
    result = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=str(value))
        cell.data_type = 's'
        result.append(cell)
    return result

def readable_cell(sheet, value, *, header=False):
    if value is None:
        text = ''
    elif isinstance(value, (dict, list)):
        text = safe_json(value)
    else:
        text = str(value)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '\ufffd', text)
    if len(text) > 30000:
        text = text[:29900] + '… [전체 원문: 데이터 / _overflow]'
    cell = WriteOnlyCell(sheet, value=text)
    cell.data_type = 's'
    if header:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='193A52')
    if len(text) <= 2048:
        try:
            parts = urlsplit(text)
            if parts.scheme in ('http', 'https') and parts.hostname:
                cell.hyperlink = text
                cell.style = 'Hyperlink'
        except ValueError:
            pass
    return cell


def export_rows(rows, output, *, mode='normal', sheet='Data', context=None):
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    kind = target.suffix.lower()
    if kind == '.jsonl':
        with target.open('w', encoding='utf-8') as stream:
            for row in rows:
                stream.write(safe_json(row) + '\n')
    elif kind in ('.csv', '.xlsx'):
        columns = columns_for(rows)
        if kind == '.csv':
            with target.open('w', newline='', encoding='utf-8-sig') as stream:
                writer = csv.writer(stream)
                writer.writerow([column[0] for column in columns])
                for row in rows:
                    writer.writerow([safe_json(resolve(row, path)) for _, path in columns])
            target.with_suffix('.csv.columns.json').write_text(canonical(columns), encoding='utf-8')
        else:
            if mode not in ('normal', 'write-only'):
                raise CollectorError('xlsx_mode must be normal or write-only')
            book = Workbook(write_only=mode == 'write-only')
            if mode == 'normal':
                book.remove(book.active)
            data = book.create_sheet(sheet_name(sheet))
            overflow = book.create_sheet('_overflow')
            mapping = book.create_sheet('_columns')
            data.append(string_cells(data, [name for name, _ in columns]))
            overflow.append(string_cells(overflow, ['reference', 'part', 'json_chunk']))
            mapping.append(string_cells(mapping, ['header', 'path_json']))
            for name, path in columns:
                mapping.append(string_cells(mapping, [name, safe_json(path)]))
            for row_index, row in enumerate(rows, 2):
                values = []
                for col_index, (_, path) in enumerate(columns, 1):
                    value = safe_json(resolve(row, path))
                    if len(value) > 30000:
                        reference = f'{row_index}:{col_index}'
                        for part, start in enumerate(range(0, len(value), 30000)):
                            overflow.append(string_cells(overflow, [reference, part, value[start:start+30000]]))
                        value = OVERFLOW + reference
                    values.append(value)
                data.append(string_cells(data, values))
            overview = book.create_sheet('Overview')
            metadata = book.create_sheet('Collection')
            names = list(dict.fromkeys(name for row in rows for name in row['fields']))
            overview.freeze_panes = 'A2'
            from openpyxl.utils import get_column_letter
            for index in range(1, len(names)+2):
                overview.column_dimensions[get_column_letter(index)].width = 35
            overview.auto_filter.ref = f'A1:{get_column_letter(len(names)+1)}{len(rows)+1}'
            overview.append([readable_cell(overview, name, header=True) for name in ['수집 페이지', *names]])
            for row_index, row in enumerate(rows, 2):
                values = [row['source_url'], *[row['fields'].get(name, {}).get('normalized') for name in names]]
                cells = [readable_cell(overview, value) for value in values]
                for column_index, cell in enumerate(cells, 1):
                    if cell.hyperlink:
                        cell.hyperlink.ref = f'{get_column_letter(column_index)}{row_index}'
                overview.append(cells)
            scope = context or {'status': None, 'partial': None, 'scope_reason': 'collection_context_not_provided'}
            scope = {**scope, 'exported_at': datetime.now(timezone.utc).isoformat(), 'exported_rows': len(rows),
                     'overview_note': '보기는 정규화값. 빈칸의 사유와 원문은 데이터 시트, 긴 값은 _overflow에 보존.'}
            metadata.column_dimensions['A'].width = 28
            metadata.column_dimensions['B'].width = 100
            for key, value in scope.items():
                metadata.append(string_cells(metadata, [key, safe_json(value)]))
            book.active = book.sheetnames.index('Overview')
            book.save(target)
    else:
        raise CollectorError('output extension must be .jsonl, .csv or .xlsx')
    return {'path': str(target), 'rows': len(rows), 'sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
            'cell_encoding': 'JSON values; XLSX long values use _overflow chunks' if kind != '.jsonl' else 'JSONL', 'xlsx_mode': mode if kind == '.xlsx' else None}

def read_xlsx(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        columns = [(row[0], json.loads(row[1])) for row in book['_columns'].iter_rows(min_row=2, values_only=True)]
        headers = tuple(name for name, _ in columns)
        data = next((book[name] for name in book.sheetnames if name not in ('_overflow', '_columns')
                     and next(book[name].iter_rows(values_only=True), ()) == headers), None)
        if data is None:
            raise CollectorError('data_sheet_headers_do_not_match_columns')
        chunks = {}
        for reference, part, value in book['_overflow'].iter_rows(min_row=2, values_only=True):
            chunks.setdefault(reference, []).append((int(part), value))
        rows = []
        for values in data.iter_rows(min_row=2, values_only=True):
            row = {'fields': {}}
            for (_, route), value in zip(columns, values):
                if value is None:
                    raise CollectorError('missing spreadsheet cell')
                if value.startswith(OVERFLOW):
                    reference = value[len(OVERFLOW):]
                    parts = sorted(chunks[reference])
                    if [p[0] for p in parts] != list(range(len(parts))):
                        raise CollectorError('overflow chunk sequence mismatch')
                    value = ''.join(p[1] for p in parts)
                decoded = json.loads(value)
                cursor = row
                for key in route[:-1]:
                    cursor = cursor.setdefault(key, {})
                cursor[route[-1]] = decoded
            rows.append(row)
        return rows
    finally:
        book.close()

def export_job(job, output, **kwargs):
    with job_lock(job):
        rows = load_records(job)
        with sqlite3.connect(f'file:{Path(job).resolve()}?mode=ro', uri=True) as db:
            source = status(db)
            config = get_meta(db, 'config')
            last_request = get_meta(db, 'last_request')
    result = export_rows(rows, output, context={**source, 'query': config, 'last_request_epoch': last_request}, **kwargs)
    result.update({'input_status': source['status'], 'input_reason': source['reason'],
                   'partial': source['partial'], 'excluded_missing_key_count': source['missing_key_count'],
                   'rejected_source': 'job SQLite rejected table' if source['missing_key_count'] else None})
    return result
