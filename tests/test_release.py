from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from sionlab_web_collector.core import collect, load_records
from sionlab_web_collector.exporting import export_job, export_rows, read_xlsx


class Fixture:
    def fetch(self, url):
        return '<article><a href="/detail/1">=SUM(1,2)</a><p>00123</p></article><a class="next" href="/page2">Next</a>', url


class ReleaseTests(unittest.TestCase):
    def test_legacy_overview_named_data_sheet_still_restores(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'legacy.xlsx'
            rows = [{'identity':'1','source_url':'https://example.invalid','fields':{}}]
            export_rows(rows,path)
            book=load_workbook(path)
            del book['Overview']; del book['Collection']
            book['Data'].title='Overview';book.save(path);book.close()
            self.assertEqual(read_xlsx(path),rows)

    def test_readable_view_context_and_roundtrip_in_both_modes(self):
        with tempfile.TemporaryDirectory() as tmp:
            job = Path(tmp)/'job.sqlite'
            config = {'start_url':'https://example.invalid/page1', 'item_selector':'article',
                'fields':{'title':{'selector':'a'}, 'url':{'selector':'a','attribute':'href','normalize':'url'},
                          'id':{'selector':'p'}}, 'unique_key':'id', 'next_selector':'a.next', 'max_pages':1}
            collect(job, config, adapter=Fixture())
            for mode in ('normal','write-only'):
                path=Path(tmp)/(mode+'.xlsx')
                exported=export_job(job,path,mode=mode,sheet='Overview')
                self.assertTrue(exported['partial'])
                self.assertEqual(read_xlsx(path),load_records(job))
                book=load_workbook(path)
                self.assertEqual(book.active.title,'Overview')
                values=list(book['Overview'].values)
                self.assertIn('title',values[0])
                self.assertIn('=SUM(1,2)',values[1])
                self.assertIn('00123',values[1])
                links=[c.hyperlink.target for c in book['Overview'][2] if c.hyperlink]
                self.assertIn('https://example.invalid/detail/1',links)
                self.assertIn('Collection',book.sheetnames)
                self.assertIn('run_limit',str(list(book['Collection'].values)))
                self.assertIn('https://example.invalid/page2',str(list(book['Collection'].values)))
                self.assertTrue(all(c.data_type!='f' for s in book for r in s for c in r))
                book.close()
