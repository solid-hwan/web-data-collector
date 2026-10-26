import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import patch

from sionlab_web_collector.core import (Blocked, CollectorError, HTTPAdapter, PlaywrightAdapter, canonical,
    check_block, collect, connect, get_meta, job_lock, load_records, parse_page, set_meta, validate_config, verify_job)
from sionlab_web_collector.exporting import export_job, export_rows, read_xlsx

ROOT = Path(__file__).resolve().parents[1]

class Handler(BaseHTTPRequestHandler):
    paths = []
    def do_GET(self):
        self.paths.append(self.path)
        status = int(self.path[1:]) if self.path in ('/401', '/403', '/429', '/500') else 200
        body = '<article><b>1</b><span>  123.40 </span></article><article><b>2</b></article><a class="next" href="/two">Next</a>'
        if self.path == '/redirect':
            self.send_response(302)
            self.send_header('Location', '/one')
            self.end_headers()
            return
        if self.path.startswith('/repeat'):
            number = int(self.path.replace('/repeat', '') or '0')
            body = '<article><b>repeat</b></article><a class="next" href="/repeat' + str(number+1) + '">Next</a>'
        elif self.path == '/two':
            body = '<article><b>2</b><span>200</span></article><article><b>3</b><span>300</span></article>'
        elif self.path == '/captcha':
            body = '<div class="g-recaptcha">challenge</div>'
        elif self.path == '/empty':
            body = '<p>Wrong selector, not verified zero records</p>'
        elif self.path == '/cycle':
            body = '<article><b>1</b></article><a class="next" href="/cycle">Next</a>'
        elif self.path == '/missingkey':
            body = '<article><span>123</span></article>'
        self.send_response(status)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.end_headers()
        self.wfile.write(body.encode())
    def log_message(self, *_):
        pass

class CollectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f'http://127.0.0.1:{cls.server.server_port}'
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.job = Path(self.temp.name) / 'job.sqlite'
        self.config = {'start_url': self.url + '/one', 'item_selector': 'article',
            'fields': {'id': {'selector': 'b'}, 'price': {'selector': 'span', 'normalize': 'decimal'}},
            'next_selector': 'a.next', 'unique_key': 'id'}
    def run_collect(self, **kwargs):
        return collect(self.job, self.config, sleep=lambda _: None, **kwargs)
    def test_pages_duplicates_missing_normalization(self):
        result = self.run_collect()
        self.assertEqual((result['status'], result['read_count'], result['stored_count'], result['duplicate_count']), ('complete', 4, 3, 1))
        rows = load_records(self.job)
        self.assertEqual(rows[0]['fields']['price'], {'raw': '  123.40 ', 'normalized': '123.40', 'missing_reason': None})
        self.assertIsNone(rows[1]['fields']['price']['normalized'])
        self.assertTrue(verify_job(self.job)['valid'])
        self.assertEqual(self.run_collect(resume=True)['stored_count'], 3)
    def test_midpage_queue_then_resume_and_config_guard(self):
        self.config['max_items'] = 1
        result = self.run_collect()
        self.assertEqual((result['pending_items'], result['completed_pages']), (1, 0))
        before = len(Handler.paths)
        result = self.run_collect(resume=True)
        self.assertEqual(len(Handler.paths), before)  # no re-fetch of partly consumed page
        self.assertEqual(result['stored_count'], 2)
        for _ in range(2):
            result = self.run_collect(resume=True)
        self.assertEqual((result['stored_count'], result['status']), (3, 'complete'))
        changed = {**self.config, 'start_url': self.url + '/other-agency'}
        with self.assertRaisesRegex(CollectorError, 'config_mismatch'):
            collect(self.job, changed, resume=True)
        with self.assertRaisesRegex(CollectorError, 'job_exists'):
            self.run_collect()
    def test_page_limit_resume(self):
        self.config['max_pages'] = 1
        first = self.run_collect()
        self.assertEqual(first['status'], 'partial')
        self.assertEqual(first['completed_pages'], 1)
        self.assertEqual(self.run_collect(resume=True)['status'], 'complete')
    def test_blocks_and_server_failure_keep_position(self):
        for path, expected in [('401', 'blocked'), ('403', 'blocked'), ('429', 'blocked'), ('captcha', 'blocked'), ('500', 'failed'), ('empty', 'failed')]:
            with self.subTest(path=path):
                self.job = Path(self.temp.name) / (path + '.sqlite')
                self.config['start_url'] = self.url + '/' + path
                result = self.run_collect()
                self.assertEqual(result['status'], expected)
                self.assertIsNone(result['total_count'])
                self.assertEqual(result['next_url'], self.config['start_url'])
                self.assertEqual(result['completed_pages'], 0)
                self.assertTrue(result['reason'])
    def test_cycle_is_failure_not_complete(self):
        self.config['start_url'] = self.url + '/cycle'
        result = self.run_collect()
        self.assertEqual(result['reason'], 'pagination_cycle')
        self.assertTrue(result['partial'])
    def test_repeated_content_under_distinct_urls_stops(self):
        self.config['start_url'] = self.url + '/repeat0'
        result = self.run_collect()
        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['reason'], 'repeated_page_content')
        self.assertEqual(result['stored_count'], 1)
    def test_redirect_does_not_issue_unbudgeted_second_request(self):
        self.config['start_url'] = self.url + '/redirect'
        before = len(Handler.paths)
        result = self.run_collect()
        self.assertEqual(result['reason'], 'redirect_requires_review')
        self.assertEqual(len(Handler.paths), before + 1)
    def test_missing_key_is_counted_and_preserved(self):
        self.config['start_url'] = self.url + '/missingkey'
        result = self.run_collect()
        self.assertEqual(result['missing_key_count'], 1)
        self.assertTrue(result['partial'])
        self.assertEqual(result['reason'], 'missing_unique_keys')
        exported = export_job(self.job, Path(self.temp.name) / 'rejected.xlsx')
        self.assertEqual(exported['excluded_missing_key_count'], 1)
        self.assertEqual(exported['input_reason'], 'missing_unique_keys')
        self.assertTrue(exported['partial'])
        with sqlite3.connect(self.job) as db:
            rejected = json.loads(db.execute('SELECT payload FROM rejected').fetchone()[0])
        self.assertEqual(rejected['fields']['price']['raw'], '123')
        self.assertTrue(verify_job(self.job)['valid'])
    def test_rate_validation_and_wait(self):
        for override in ({'concurrency': 2}, {'requests_per_second': 2}, {'max_items': 0}, {'unique_key': 'absent'}, {'item_selector': '['}):
            with self.assertRaises(CollectorError):
                validate_config({**self.config, **override})
        waits = []
        collect(self.job, self.config, sleep=waits.append)
        self.assertEqual(len(waits), 2)
        self.assertGreater(waits[1], 0)
        self.assertLessEqual(waits[1], 1)
    def test_concurrent_process_is_rejected(self):
        self.run_collect()
        with job_lock(self.job):
            result = subprocess.run([sys.executable, '-m', 'sionlab_web_collector', 'resume', '--job', str(self.job)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn('job_busy', result.stdout)
    def test_real_process_crash_before_and_after_commit(self):
        config = Path(self.temp.name) / 'config.json'
        config.write_text(json.dumps(self.config))
        child = '''import json,os,sys
from sionlab_web_collector.core import collect
config=json.load(open(sys.argv[1]))
def hook(phase):
 if phase==sys.argv[3]: os._exit(91)
collect(sys.argv[2],config,checkpoint_hook=hook,sleep=lambda _:None)
'''
        for phase, expected_count in [('before_commit', 0), ('after_commit', 2)]:
            with self.subTest(phase=phase):
                self.job = Path(self.temp.name) / (phase + '.sqlite')
                result = subprocess.run([sys.executable, '-c', child, str(config), str(self.job), phase], capture_output=True)
                self.assertEqual(result.returncode, 91)
                self.assertEqual(len(load_records(self.job)), expected_count)
                resumed = self.run_collect(resume=True)
                self.assertEqual((resumed['stored_count'], resumed['read_count'], resumed['duplicate_count']), (3, 4, 1))
                self.assertTrue(verify_job(self.job)['valid'])
    def test_cli_end_to_end(self):
        config = Path(self.temp.name) / 'config.json'
        config.write_text(json.dumps(self.config))
        workbook = Path(self.temp.name) / 'result.xlsx'
        commands = [['crawl', '--job', str(self.job), '--config', str(config)],
                    ['resume', '--job', str(self.job)],
                    ['export', '--job', str(self.job), '--output', str(workbook), '--xlsx-mode', 'write-only'],
                    ['verify', '--job', str(self.job), '--xlsx', str(workbook)]]
        for command in commands:
            result = subprocess.run([sys.executable, '-m', 'sionlab_web_collector', *command], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            payload = json.loads(result.stdout)
        self.assertTrue(payload['xlsx_matches'])
    def test_dynamic_adapter_contract_without_browser_install(self):
        from types import SimpleNamespace
        class Browser:
            closed = False
            def new_context(self, **kwargs): return self
            def route(self, pattern, handler): self.handler = handler
            def new_page(self): return self
            def goto(self, url, **kwargs):
                self.url=url
                self.fetch_calls=[]
                def fetch(**options):
                    self.fetch_calls.append(options)
                    return SimpleNamespace(status=getattr(self, 'response_status', 200))
                route=SimpleNamespace(request=SimpleNamespace(method='GET', resource_type='document'), fetch=fetch, fulfill=lambda **kwargs: None, abort=lambda: None)
                self.handler(route)
                self.handler(route)
                return SimpleNamespace(status=200)
            def wait_for_selector(self, selector, **kwargs): self.waited = (selector, kwargs)
            def content(self): return '<article><b>9</b><span>42</span></article>'
            def close(self): self.closed=True
        browser = Browser()
        class Engine:
            def __enter__(self): return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kwargs: browser))
            def __exit__(self, *_): pass
        stub = SimpleNamespace(sync_playwright=Engine)
        with patch.dict(sys.modules, {'playwright.sync_api': stub}), patch('sionlab_web_collector.core.time.sleep') as sleeper:
            html, url = PlaywrightAdapter(wait_selector='article').fetch(self.url)
        self.assertEqual(browser.waited[0], 'article')
        self.assertEqual(browser.waited[1]['state'], 'attached')
        self.assertEqual(len(browser.fetch_calls), 2)
        self.assertTrue(all(call['max_redirects'] == 0 and call['max_retries'] == 0 for call in browser.fetch_calls))
        self.assertGreater(sleeper.call_args_list[1].args[0], 0)
        browser.response_status = 403
        with patch.dict(sys.modules, {'playwright.sync_api': stub}), patch('sionlab_web_collector.core.time.sleep'):
            with self.assertRaises(Blocked):
                PlaywrightAdapter().fetch(self.url)
        self.assertEqual(len(browser.fetch_calls), 1)
        self.assertIn('<b>9</b>', html)
        self.assertTrue(browser.closed)
        self.assertEqual(url, self.url)
    def test_nested_invalid_and_cross_origin(self):
        config = {**self.config, 'fields': {'id': {'selector': 'b'}, 'nested': {'selector': 'i', 'normalize': 'json'}}}
        rows, _ = parse_page('<article><b>1</b><i>{"a":[1,2]}</i></article>', self.url, config)
        self.assertEqual(rows[0]['fields']['nested']['normalized'], {'a': [1,2]})
        rows, _ = parse_page('<article><b>1</b><i>no</i></article>', self.url, config)
        self.assertEqual(rows[0]['fields']['nested']['missing_reason'], 'normalization_failed')
        with self.assertRaisesRegex(CollectorError, 'cross_origin'):
            parse_page('<a class="next" href="https://elsewhere.invalid/">Next</a>', self.url, self.config)
    def test_login_challenge_is_blocked(self):
        for html in ['<input type="password">', '<h1>Verify you are human</h1>']:
            with self.assertRaises(Blocked): check_block(200, html)

class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
    def records(self, count):
        return [{'identity': str(i), 'source_url': 'https://example.invalid/' + str(i), 'fields': {
            '금액': {'raw': '900719925474099312345.67', 'normalized': '900719925474099312345.67', 'missing_reason': None},
            '=formula': {'raw': '=HYPERLINK("bad")', 'normalized': {'nested': [i, None, '한글', '\ufffe', '\ud800']}, 'missing_reason': None},
            'a/b': {'raw': 'x\x00\x01' + ('한글' * 20000 if i == 0 else str(i)), 'normalized': None, 'missing_reason': 'example'},
            'a?b': {'raw': '@SUM(1)', 'normalized': '+1', 'missing_reason': None}}} for i in range(count)]
    def test_two_xlsx_modes_20000_rows_exact_roundtrip(self):
        from openpyxl import load_workbook
        rows = self.records(20000)
        for mode in ('normal', 'write-only'):
            with self.subTest(mode=mode):
                target = Path(self.temp.name) / (mode + '.xlsx')
                result = export_rows(rows, target, mode=mode, sheet='_overflow')
                self.assertEqual(result['rows'], 20000)
                self.assertEqual(read_xlsx(target), rows)
                book = load_workbook(target, read_only=True, data_only=False)
                self.assertEqual(book.sheetnames, ['_overflow_data', '_overflow', '_columns', 'Overview', 'Collection'])
                for row in book['_overflow_data'].iter_rows(min_row=2, max_row=2):
                    self.assertTrue(all(cell.data_type == 's' for cell in row))
                book.close()
    def test_jsonl_and_csv_exact_cells_and_header_mapping(self):
        import csv
        rows = self.records(2)
        for suffix in ('.jsonl', '.csv'):
            target = Path(self.temp.name) / ('result' + suffix)
            export_rows(rows, target)
            if suffix == '.jsonl':
                self.assertEqual([json.loads(x) for x in target.read_text().splitlines()], rows)
            else:
                csv.field_size_limit(10_000_000)
                with target.open(encoding='utf-8-sig', newline='') as stream:
                    values = list(csv.reader(stream))
                self.assertEqual(len(values), 3)
                self.assertEqual(len(values[0]), len(set(values[0])))
                self.assertTrue(all(not cell.startswith(('=', '+', '-', '@')) for cell in values[1]))
                self.assertEqual(json.loads(values[1][2]), rows[0]['fields']['금액']['raw'])
                self.assertTrue(target.with_suffix('.csv.columns.json').is_file())
    def test_korean_cells_are_human_readable(self):
        from openpyxl import load_workbook
        rows = self.records(1)
        target = Path(self.temp.name) / 'korean.xlsx'
        export_rows(rows, target)
        book = load_workbook(target, read_only=True)
        self.assertIn('한글', book['Data']['G2'].value)
        book.close()
        self.assertEqual(read_xlsx(target), rows)
    def test_empty_export_is_valid(self):
        target = Path(self.temp.name) / 'empty.xlsx'
        export_rows([], target)
        self.assertEqual(read_xlsx(target), [])

if __name__ == '__main__':
    unittest.main()
